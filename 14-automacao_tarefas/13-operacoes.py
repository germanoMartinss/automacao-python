from openpyxl import load_workbook
from openpyxl.drawing.image import Image

# 1 - Lendo planilha Excel
wb = load_workbook(filename='files/gastos.xlsx')
planilha = wb['Sheet']

valor_total = 0
for i in range(2, 10):
    valor = int(planilha['B%s' %i].value)
    valor_total += valor
# print("O valor total foi de {0}".format(valor_total))
planilha['B10'] = valor_total

# wb.save(filename='files/gastos.xlsx')

# 2 - Mesclar Céluas
planilha['A11'] = "Total"
planilha.merge_cells('A11:B11')
# wb.save(filename='files/gastos.xlsx')

# 3 - Inserir Imagem
img = Image('files/bb_preco.png')
planilha.add_image(img, 'A13')
wb.save(filename='files/gastos.xlsx')



