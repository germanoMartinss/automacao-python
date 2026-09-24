from openpyxl import load_workbook

arquivo = 'files/teste.xlsx'
wb = load_workbook(arquivo)
ws = wb.active
print(ws['A1'].value)